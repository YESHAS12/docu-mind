"""Document loaders for various file types (PDF, PPTX, XLSX, TXT, MD).
Each loader returns a list of LangChain Document objects with standardized metadata.
"""

from pathlib import Path
from typing import List
from langchain_core.documents import Document
import pypdf
from pptx import Presentation
import pandas as pd


def load_pdf(file_path: Path) -> List[Document]:
    """Extract text from PDF pages using pypdf."""
    documents = []
    reader = pypdf.PdfReader(str(file_path))
    for idx, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        text = text.strip()
        if text:
            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": file_path.name,
                        "type": "pdf",
                        "page": idx + 1
                    }
                )
            )
    return documents


def load_pptx(file_path: Path) -> List[Document]:
    """Extract text from PPTX slides using python-pptx."""
    documents = []
    prs = Presentation(str(file_path))
    for idx, slide in enumerate(prs.slides):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_texts.append(line)
        if slide_texts:
            documents.append(
                Document(
                    page_content="\n".join(slide_texts),
                    metadata={
                        "source": file_path.name,
                        "type": "pptx",
                        "slide": idx + 1
                    }
                )
            )
    return documents


def load_xlsx(file_path: Path) -> List[Document]:
    """Extract tabular data from XLSX into textual representations for semantic indexing."""
    documents = []
    excel_file = pd.ExcelFile(str(file_path))
    for sheet_name in excel_file.sheet_names:
        df = pd.read_excel(excel_file, sheet_name=sheet_name)
        # Create a structured text summary and row-by-row descriptions
        rows_text = []
        rows_text.append(f"Sheet '{sheet_name}' with columns: {', '.join(df.columns)}")
        for _, row in df.iterrows():
            row_repr = ", ".join([f"{col}: {val}" for col, val in row.items()])
            rows_text.append(f"Record: {row_repr}")

        page_content = "\n".join(rows_text)
        documents.append(
            Document(
                page_content=page_content,
                metadata={
                    "source": file_path.name,
                    "type": "xlsx",
                    "sheet": sheet_name
                }
            )
        )
    return documents


def load_text_or_markdown(file_path: Path) -> List[Document]:
    """Read plain text or markdown file."""
    text = file_path.read_text(encoding="utf-8").strip()
    if not text:
        return []
    filetype = "markdown" if file_path.suffix.lower() == ".md" else "text"
    return [
        Document(
            page_content=text,
            metadata={
                "source": file_path.name,
                "type": filetype
            }
        )
    ]


def load_document(file_path: Path) -> List[Document]:
    """Dispatch to the appropriate loader based on file extension."""
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return load_pdf(file_path)
    elif suffix in (".pptx", ".ppt"):
        return load_pptx(file_path)
    elif suffix in (".xlsx", ".xls"):
        return load_xlsx(file_path)
    elif suffix in (".md", ".txt"):
        return load_text_or_markdown(file_path)
    else:
        print(f"Skipping unsupported file extension: {suffix} ({file_path.name})")
        return []


def load_directory(directory_path: Path) -> List[Document]:
    """Load all supported documents from a directory."""
    all_docs = []
    for file_path in sorted(directory_path.iterdir()):
        if file_path.is_file():
            docs = load_document(file_path)
            all_docs.extend(docs)
    return all_docs
