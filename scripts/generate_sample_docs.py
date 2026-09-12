"""Generate sample documents for DocuMind demonstration.
Creates realistic PDF, PPTX, XLSX, and Markdown files in sample_docs/.
"""

from pathlib import Path
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib import colors


def create_pdf(file_path: Path, title: str, sections: list[tuple[str, list[str]]]):
    """Helper to generate a formatted PDF with title and sections."""
    c = canvas.Canvas(str(file_path), pagesize=letter)
    width, height = letter

    # Title
    c.setFont("Helvetica-Bold", 18)
    c.setFillColor(colors.HexColor("#1A365D"))
    c.drawString(54, height - 54, title)

    c.setStrokeColor(colors.HexColor("#CBD5E1"))
    c.setLineWidth(1)
    c.line(54, height - 64, width - 54, height - 64)

    y = height - 90
    for section_title, bullets in sections:
        if y < 100:
            c.showPage()
            y = height - 54

        c.setFont("Helvetica-Bold", 13)
        c.setFillColor(colors.HexColor("#2B6CB0"))
        c.drawString(54, y, section_title)
        y -= 20

        c.setFont("Helvetica", 10)
        c.setFillColor(colors.HexColor("#2D3748"))
        for bullet in bullets:
            if y < 60:
                c.showPage()
                y = height - 54
            c.drawString(68, y, f"- {bullet}")
            y -= 16
        y -= 12

    c.save()


def generate_company_policy_pdf(output_dir: Path):
    path = output_dir / "company_policy.pdf"
    sections = [
        ("1. Remote Work & Flexible Hours", [
            "Employees are eligible for a hybrid work model allowing up to 3 days of work-from-home per week.",
            "Core collaborative working hours are 10:00 AM to 4:00 PM EST.",
            "Employees working remotely must ensure a secure high-speed internet connection and VPN compliance."
        ]),
        ("2. Paid Time Off (PTO) & Leave", [
            "Full-time team members receive 25 days of paid annual leave per calendar year.",
            "Up to 5 unused annual leave days may be carried over into the following fiscal year.",
            "Sick leave allowance is 10 days annually without requiring medical certification for under 3 consecutive days.",
            "Each team member receives 5 paid personal development and learning days per year."
        ]),
        ("3. Equipment & Home Office Stipend", [
            "All new hires receive a one-time home office setup reimbursement of up to $1,500.",
            "A recurring monthly internet and connectivity stipend of $100 is disbursed with payroll.",
            "Standard laptop refresh cycle occurs every 24 months for engineering and 36 months for general staff."
        ]),
        ("4. Travel & Corporate Expenses", [
            "Daily meal expenses during business travel are capped at $75 per day with receipts.",
            "All domestic flight bookings must be economy class unless flight duration exceeds 5 continuous hours."
        ])
    ]
    create_pdf(path, "DocuMind Corp - Global Employee Handbook & Workplace Policy", sections)
    print(f"Generated {path}")


def generate_cloud_architecture_pdf(output_dir: Path):
    path = output_dir / "cloud_architecture.pdf"
    sections = [
        ("1. System Overview & Architecture", [
            "The platform is deployed across multiple cloud availability zones using Kubernetes (EKS/GKE).",
            "Core microservices communicate asynchronously via Apache Kafka message brokers.",
            "API Gateway manages rate limiting, authentication, and ingress routing."
        ]),
        ("2. Reliability & High Availability SLAs", [
            "The infrastructure is designed for 99.99% annual uptime service level agreement (SLA).",
            "Recovery Time Objective (RTO) is strictly under 15 minutes in disaster recovery scenarios.",
            "Recovery Point Objective (RPO) is under 1 minute via synchronous cross-region transaction logs."
        ]),
        ("3. Performance & Latency Budgets", [
            "P95 API response latency target is under 120 milliseconds for cached endpoints.",
            "P99 API response latency budget is capped at 250 milliseconds under peak concurrent load of 50,000 RPS.",
            "Vector database semantic search latency SLA is under 45 milliseconds for top-k retrieval."
        ]),
        ("4. Data Persistence & Security", [
            "Primary relational data resides in PostgreSQL with multi-AZ replication and automated failover.",
            "Automated point-in-time recovery (PITR) backups are maintained for a 30-day retention window.",
            "All data is encrypted in transit using TLS 1.3 and at rest using AES-256 keys managed via KMS."
        ])
    ]
    create_pdf(path, "DocuMind Cloud Architecture & Infrastructure Specification", sections)
    print(f"Generated {path}")


def generate_presentation_pptx(output_dir: Path):
    path = output_dir / "quarterly_presentation.pptx"
    prs = Presentation()

    # Slide 1: Title
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "DocuMind Q3 Strategic Business Review"
    subtitle.text = "Quarterly Highlights, Financial Performance, and Product Vision"

    # Slide 2: Achievements
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Q3 Key Milestones & Growth"
    body = slide.shapes.placeholders[1]
    tf = body.text_frame
    tf.text = "Surpassed 15,000 active monthly enterprise users (38% YoY growth)."
    p = tf.add_paragraph()
    p.text = "Maintained 99.98% platform operational availability throughout the quarter."
    p = tf.add_paragraph()
    p.text = "Successfully rolled out European data sovereign region complying with GDPR standards."
    p = tf.add_paragraph()
    p.text = "Reduced average document vector indexing latency from 4.2 seconds to 1.6 seconds."

    # Slide 3: Roadmap
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Strategic Objectives for Q4"
    body = slide.shapes.placeholders[1]
    tf = body.text_frame
    tf.text = "Deploy multi-agent LangGraph workflow for document reasoning and verification."
    p = tf.add_paragraph()
    p.text = "Onboard 50 new enterprise partners across financial services and legal sectors."
    p = tf.add_paragraph()
    p.text = "Launch native Model Context Protocol (MCP) server integration for client tooling."
    p = tf.add_paragraph()
    p.text = "Target annualized recurring revenue (ARR) milestone of $22 Million by year-end."

    prs.save(str(path))
    print(f"Generated {path}")


def generate_financial_xlsx(output_dir: Path):
    path = output_dir / "financial_q3.xlsx"
    data = [
        {"Department": "Engineering", "Budget": 500000, "Actual_Spend": 480000, "Revenue_Contribution": 1200000, "Headcount": 35},
        {"Department": "Sales", "Budget": 350000, "Actual_Spend": 365000, "Revenue_Contribution": 2800000, "Headcount": 22},
        {"Department": "Marketing", "Budget": 200000, "Actual_Spend": 195000, "Revenue_Contribution": 650000, "Headcount": 12},
        {"Department": "Operations", "Budget": 150000, "Actual_Spend": 140000, "Revenue_Contribution": 300000, "Headcount": 8},
        {"Department": "Customer_Success", "Budget": 120000, "Actual_Spend": 115000, "Revenue_Contribution": 450000, "Headcount": 14},
    ]
    df = pd.DataFrame(data)

    with pd.ExcelWriter(str(path), engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Q3_Summary", index=False)

    print(f"Generated {path}")


def generate_markdown(output_dir: Path):
    path = output_dir / "team_handbook.md"
    content = """# Engineering Team Handbook & Operating Principles

Welcome to the DocuMind Engineering organization. This guide establishes our day-to-day engineering practices and development culture.

## 1. Code Review & Quality Standards
- Every Pull Request (PR) requires a minimum of **2 peer approvals** before merging to `main`.
- All automated continuous integration (CI) tests, security vulnerability scans, and linters must pass with zero errors.
- PR titles must follow Conventional Commits standard (e.g., `feat:`, `fix:`, `refactor:`, `docs:`).

## 2. Release & Deployment Cadence
- Deployments to the staging environment occur continuously upon merging to the default branch.
- Production releases are scheduled on **Tuesday and Thursday mornings at 10:00 AM EST** to maximize team availability.
- No production deployments are allowed on Fridays, weekends, or public holidays except for critical emergency security hotfixes.

## 3. Incident Management & On-Call
- P1 critical incidents require an on-call response time within **15 minutes**.
- A blameless post-mortem root cause analysis (RCA) must be published within **48 hours** of incident resolution.
- On-call rotations run weekly from Monday 9:00 AM EST to the following Monday 9:00 AM EST.
"""
    path.write_text(content, encoding="utf-8")
    print(f"Generated {path}")


def main():
    output_dir = Path(__file__).resolve().parent.parent / "sample_docs"
    output_dir.mkdir(parents=True, exist_ok=True)

    generate_company_policy_pdf(output_dir)
    generate_cloud_architecture_pdf(output_dir)
    generate_presentation_pptx(output_dir)
    generate_financial_xlsx(output_dir)
    generate_markdown(output_dir)
    print("All sample documents successfully generated in sample_docs/!")


if __name__ == "__main__":
    main()
