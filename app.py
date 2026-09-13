"""DocuMind — Multi-Agent Document Intelligence Assistant
Streamlit Entrypoint Application.
Supports both local execution and zero-cost deployment on Streamlit Community Cloud.
"""

import sys
import os
import time
from pathlib import Path
import streamlit as st
import pandas as pd

# Ensure root directory in sys.path
BASE_DIR = Path(__file__).resolve().parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from src.config import get_groq_api_key, DEFAULT_GROQ_MODEL, SAMPLE_DOCS_DIR
from src.retrieval.vectorstore import vectorstore_exists, add_file_to_vectorstore
from src.scripts_runner import ensure_knowledge_base_ready
from src.agents.graph import run_agent

# Configure Streamlit page
st.set_page_config(
    page_title="DocuMind — Document Intelligence",
    page_icon="🧠",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Custom Design & Glassmorphism Styling
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Outfit:wght@300;400;500;600;700&family=Inter:wght@400;500;600&display=swap');

html, body, [class*="css"] {
    font-family: 'Inter', -apple-system, BlinkMacSystemFont, sans-serif;
}

h1, h2, h3, .brand-title {
    font-family: 'Outfit', sans-serif !important;
    font-weight: 700;
}

/* App Header Banner */
.hero-header {
    background: linear-gradient(135deg, rgba(30, 41, 59, 0.7) 0%, rgba(15, 23, 42, 0.9) 100%);
    border: 1px solid rgba(255, 255, 255, 0.1);
    border-radius: 16px;
    padding: 24px 30px;
    margin-bottom: 25px;
    backdrop-filter: blur(10px);
    box-shadow: 0 8px 32px 0 rgba(0, 0, 0, 0.37);
}

.hero-title {
    font-size: 2.2rem;
    font-weight: 700;
    background: linear-gradient(90deg, #60A5FA, #A78BFA, #F472B6);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
    margin: 0;
}

.hero-subtitle {
    color: #94A3B8;
    font-size: 1.02rem;
    margin-top: 6px;
    margin-bottom: 0;
}

/* Route Badges */
.badge {
    display: inline-block;
    padding: 4px 12px;
    border-radius: 9999px;
    font-size: 0.78rem;
    font-weight: 600;
    letter-spacing: 0.5px;
    text-transform: uppercase;
    margin-bottom: 8px;
}
.badge-docs {
    background-color: rgba(59, 130, 246, 0.2);
    color: #93C5FD;
    border: 1px solid rgba(59, 130, 246, 0.4);
}
.badge-spreadsheet {
    background-color: rgba(16, 185, 129, 0.2);
    color: #6EE7B7;
    border: 1px solid rgba(16, 185, 129, 0.4);
}
.badge-general {
    background-color: rgba(168, 85, 247, 0.2);
    color: #D8B4FE;
    border: 1px solid rgba(168, 85, 247, 0.4);
}

/* Citation Chips */
.citation-chip {
    display: inline-flex;
    align-items: center;
    background: rgba(255, 255, 255, 0.05);
    border: 1px solid rgba(255, 255, 255, 0.15);
    border-radius: 8px;
    padding: 4px 10px;
    margin: 4px 6px 4px 0;
    font-size: 0.82rem;
    color: #E2E8F0;
}

/* Metric Cards */
.metric-card {
    background: rgba(30, 41, 59, 0.5);
    border: 1px solid rgba(255, 255, 255, 0.08);
    border-radius: 12px;
    padding: 14px;
    text-align: center;
    margin-bottom: 12px;
}
.metric-val {
    font-size: 1.6rem;
    font-weight: 700;
    color: #60A5FA;
}
.metric-lbl {
    font-size: 0.78rem;
    color: #94A3B8;
    text-transform: uppercase;
    letter-spacing: 0.5px;
}
</style>
""", unsafe_allow_html=True)


# Initialize Session State
if "messages" not in st.session_state:
    st.session_state.messages = [
        {
            "role": "assistant",
            "content": "Hello! I am **DocuMind**, your multi-agent document intelligence assistant.\n\nI can query across policies, technical architecture, presentation decks, and perform exact numerical calculations on financial spreadsheets. How can I help you today?",
            "route": "general",
            "sources": [],
            "trace": [],
            "retry_count": 0
        }
    ]

if "total_queries" not in st.session_state:
    st.session_state.total_queries = 0

if "total_retries" not in st.session_state:
    st.session_state.total_retries = 0

if "prompt_to_submit" not in st.session_state:
    st.session_state.prompt_to_submit = None


# Ensure vector store is initialized
ensure_knowledge_base_ready()

# Sidebar: System Metrics & Document Directory
with st.sidebar:
    st.markdown("### 🧠 DocuMind Architecture")
    st.caption("Multi-Agent LangGraph System with Self-Correcting Retry Loop & Pandas Math")

    st.markdown("---")
    st.markdown("#### 📊 Session Activity")
    col1, col2 = st.columns(2)
    with col1:
        st.markdown(f"""
        <div class="metric-card">
            <div class="metric-val">{st.session_state.total_queries}</div>
            <div class="metric-lbl">Queries</div>
        </div>
        """, unsafe_allow_html=True)
    with col2:
        avg_retries = (st.session_state.total_retries / st.session_state.total_queries) if st.session_state.total_queries > 0 else 0.0
        st.markdown(f"""
        <div class="metric-card">
            <div class="metric-val">{avg_retries:.1f}</div>
            <div class="metric-lbl">Avg Retries</div>
        </div>
        """, unsafe_allow_html=True)

    st.markdown("---")
    st.markdown("#### 📤 Upload New Documents")
    st.caption("Upload files to dynamically index them into the vector knowledge base:")
    uploaded_files = st.file_uploader(
        "Upload PDF, PPTX, XLSX, MD, or TXT",
        type=["pdf", "pptx", "xlsx", "md", "txt"],
        accept_multiple_files=True,
        label_visibility="collapsed",
        key="doc_uploader"
    )
    if uploaded_files:
        newly_indexed = 0
        for uploaded_file in uploaded_files:
            dest_file = SAMPLE_DOCS_DIR / uploaded_file.name
            if not dest_file.exists():
                with open(dest_file, "wb") as f:
                    f.write(uploaded_file.getbuffer())
                with st.spinner(f"Indexing '{uploaded_file.name}'..."):
                    chunk_count = add_file_to_vectorstore(dest_file)
                st.toast(f"✅ Indexed '{uploaded_file.name}' ({chunk_count} chunks)!", icon="🎉")
                newly_indexed += 1
        if newly_indexed > 0:
            st.rerun()

    st.markdown("---")
    st.markdown("#### 📂 Knowledge Base Documents")
    st.caption("Click any document to inspect contents or download:")
    if SAMPLE_DOCS_DIR.exists():
        for doc_file in sorted(SAMPLE_DOCS_DIR.iterdir()):
            if doc_file.is_file():
                ext = doc_file.suffix.lower()
                icon = "📄"
                if ext == ".pdf":
                    icon = "📕"
                elif ext == ".xlsx":
                    icon = "📊"
                elif ext in (".pptx", ".ppt"):
                    icon = "📽️"
                elif ext == ".md":
                    icon = "📝"

                with st.expander(f"{icon} {doc_file.name}", expanded=False):
                    try:
                        with open(doc_file, "rb") as f:
                            file_bytes = f.read()
                        st.download_button(
                            label=f"⬇️ Download {doc_file.name}",
                            data=file_bytes,
                            file_name=doc_file.name,
                            key=f"dl_{doc_file.name}",
                            use_container_width=True
                        )
                    except Exception as e:
                        st.caption(f"Download unavailable: {e}")

                    # Content preview
                    if ext in (".md", ".txt"):
                        try:
                            st.markdown(doc_file.read_text(encoding="utf-8")[:1200] + ("..." if doc_file.stat().st_size > 1200 else ""))
                        except Exception:
                            st.caption("Preview unavailable.")
                    elif ext == ".xlsx":
                        try:
                            xls = pd.ExcelFile(doc_file)
                            for sheet in xls.sheet_names[:2]:
                                st.caption(f"Sheet: `{sheet}`")
                                df_preview = pd.read_excel(xls, sheet_name=sheet)
                                st.dataframe(df_preview.head(5), use_container_width=True)
                        except Exception as e:
                            st.caption(f"Preview unavailable: {e}")
                    elif ext == ".pdf":
                        try:
                            import pypdf
                            reader = pypdf.PdfReader(str(doc_file))
                            page_text = reader.pages[0].extract_text() if reader.pages else ""
                            st.text_area("Page 1 Preview", value=page_text[:600], height=120, disabled=True)
                        except Exception:
                            st.caption("PDF preview unavailable.")
                    elif ext in (".pptx", ".ppt"):
                        try:
                            from pptx import Presentation
                            prs = Presentation(str(doc_file))
                            slide_texts = []
                            for idx, s in enumerate(prs.slides[:2], 1):
                                for shape in s.shapes:
                                    if shape.has_text_frame:
                                        slide_texts.append(f"[Slide {idx}] {shape.text_frame.text.strip()}")
                            st.text_area("Slides Preview", value="\n".join(slide_texts)[:600], height=120, disabled=True)
                        except Exception:
                            st.caption("PPTX preview unavailable.")

    st.markdown("---")
    st.markdown("#### 💡 Try Sample Questions")
    sample_queries = [
        "What is the remote work policy and leave allowance?",
        "What is the system uptime SLA and P95 latency target?",
        "What is the total actual spend and budget in Q3?",
        "Which department has the highest headcount?",
        "What are the pull request review requirements?"
    ]
    for sq in sample_queries:
        if st.button(sq, key=f"sq_{sq}", use_container_width=True):
            st.session_state.prompt_to_submit = sq

    st.markdown("---")
    if st.button("🗑️ Clear Conversation", use_container_width=True):
        st.session_state.messages = [st.session_state.messages[0]]
        st.session_state.total_queries = 0
        st.session_state.total_retries = 0
        st.rerun()


# Main Chat Interface
st.markdown("""
<div class="hero-header">
    <h1 class="hero-title">DocuMind Intelligence</h1>
    <p class="hero-subtitle">Agentic Q&A with Semantic Vector Search, Self-Correcting Grading Loops, and Precise Pandas Computation</p>
</div>
""", unsafe_allow_html=True)

# Render Chat History
for msg in st.session_state.messages:
    with st.chat_message(msg["role"], avatar="🤖" if msg["role"] == "assistant" else "👤"):
        # If assistant message has routing metadata, display badge
        if msg["role"] == "assistant" and msg.get("route"):
            route = msg["route"]
            badge_class = f"badge-{route}"
            retries = msg.get("retry_count", 0)
            retry_tag = f" • {retries} retries" if retries > 0 else ""
            st.markdown(f'<span class="badge {badge_class}">Route: {route.upper()}{retry_tag}</span>', unsafe_allow_html=True)

        st.markdown(msg["content"])

        # Display citations if available
        sources = msg.get("sources", [])
        if sources:
            with st.expander("📑 Referenced Sources & Citations", expanded=False):
                for src in sources:
                    st.markdown(f'<span class="citation-chip">📄 {src}</span>', unsafe_allow_html=True)

        # Display reasoning trace if available
        trace = msg.get("trace", [])
        if trace and len(trace) > 1:
            with st.expander("🔍 View Agent Execution Trace", expanded=False):
                for step in trace:
                    st.markdown(f"• `{step}`")


# Handle User Input (via chat input or sidebar sample button)
user_prompt = st.chat_input("Ask a question about policies, architecture, slides, or spreadsheet math...")
if st.session_state.prompt_to_submit:
    user_prompt = st.session_state.prompt_to_submit
    st.session_state.prompt_to_submit = None

if user_prompt:
    clean_prompt = user_prompt.strip()

    # Guardrail 1: Input length capping
    if len(clean_prompt) > 500:
        st.warning("⚠️ Query exceeds 500 characters. Please condense your question.")
        st.stop()

    # Guardrail 2: Basic prompt injection filter
    lower_prompt = clean_prompt.lower()
    if "ignore previous instructions" in lower_prompt or "disregard all previous" in lower_prompt:
        clean_prompt = "Hello, what documents do you have access to?"

    # Append user message
    st.session_state.messages.append({"role": "user", "content": clean_prompt})
    with st.chat_message("user", avatar="👤"):
        st.markdown(clean_prompt)

    # Process through LangGraph Agent
    with st.chat_message("assistant", avatar="🤖"):
        with st.spinner("🧠 DocuMind agents evaluating and retrieving across corpus..."):
            start_time = time.time()
            try:
                result = run_agent(clean_prompt)
                answer = result.get("answer", "No answer generated.")
                route = result.get("route", "docs")
                sources = result.get("sources", [])
                trace = result.get("trace", [])
                retry_count = result.get("retry_count", 0)

                # Update session stats
                st.session_state.total_queries += 1
                st.session_state.total_retries += retry_count

                # Display Route Badge
                badge_class = f"badge-{route}"
                retry_tag = f" • {retry_count} retries" if retry_count > 0 else ""
                st.markdown(f'<span class="badge {badge_class}">Route: {route.upper()}{retry_tag}</span>', unsafe_allow_html=True)

                st.markdown(answer)

                if sources:
                    with st.expander("📑 Referenced Sources & Citations", expanded=False):
                        for src in sources:
                            st.markdown(f'<span class="citation-chip">📄 {src}</span>', unsafe_allow_html=True)

                if trace:
                    with st.expander("🔍 View Agent Execution Trace", expanded=False):
                        for step in trace:
                            st.markdown(f"• `{step}`")

                # Append assistant message
                st.session_state.messages.append({
                    "role": "assistant",
                    "content": answer,
                    "route": route,
                    "sources": sources,
                    "trace": trace,
                    "retry_count": retry_count
                })

            except Exception as e:
                err_msg = f"An error occurred during agent execution: {str(e)}"
                st.error(err_msg)
                st.session_state.messages.append({
                    "role": "assistant",
                    "content": err_msg,
                    "route": "error",
                    "sources": [],
                    "trace": [str(e)],
                    "retry_count": 0
                })
